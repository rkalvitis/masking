"""From out/bench/iou_pairs.csv: highlighted spreadsheet (red < 0.90, orange < 0.95, thresholds
editable) and a visual sheet (original + every mask set) of the frames with big disagreement."""
import os
import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter as L
from PIL import Image
from beetlemask import common as C, viz
D = "/Users/robertskalvitis/Library/CloudStorage/OneDrive-USI/SUPSI/dataset_popillia_15_09_2026_FS"
BIG, MOD = 0.90, 0.95
df = pd.read_csv("out/bench/iou_pairs.csv"); cols = [c for c in df.columns if " vs " in c]

wb = Workbook(); ws = wb.active; ws.title = "IoU per frame"
hdr = ["image", "set", "view"] + cols; ws.append(hdr)
for _, r in df.iterrows():
    ws.append([r["image"], r["set"], r["view"]] + [round(float(r[c]), 4) if pd.notna(r[c]) else None for c in cols])
n = len(df) + 1
bold = Font(name="Arial", bold=True); reg = Font(name="Arial")
for row in ws.iter_rows(min_row=1, max_row=n, max_col=len(hdr)):
    for c in row: c.font = reg
for c in ws[1]: c.font = bold
# summary block two columns to the right of the data
S = 5 + len(cols)                      # first summary column index
lab, med, mn, big, mod = (L(S + k) for k in range(5))
thr_big, thr_mod = f"${med}$3", f"${med}$4"
for j in range(4, 4 + len(cols)):
    col = L(j)
    for i in range(2, n + 1): ws[f"{col}{i}"].number_format = "0.000"
    rng = f"{col}2:{col}{n}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=[thr_big], fill=PatternFill("solid", start_color="FFC7CE"), font=Font(name="Arial", color="9C0006", bold=True)))
    ws.conditional_formatting.add(rng, CellIsRule(operator="between", formula=[thr_big, thr_mod], fill=PatternFill("solid", start_color="FFEB9C"), font=Font(name="Arial", color="9C5700")))
    ws.column_dimensions[col].width = 19
ws.column_dimensions["A"].width = 10
ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:{L(len(hdr))}{n}"
ws[f"{lab}2"] = "Thresholds (edit the yellow cells):"; ws[f"{lab}2"].font = bold
ws[f"{lab}3"] = "big disagreement: IoU <"; ws[f"{med}3"] = BIG
ws[f"{lab}4"] = "moderate: IoU <"; ws[f"{med}4"] = MOD
for c in (f"{med}3", f"{med}4"): ws[c].font = Font(name="Arial", color="0000FF"); ws[c].fill = PatternFill("solid", start_color="FFFF00")
ws[f"{lab}5"] = "red = below the big-disagreement threshold, orange = between the two"; ws[f"{lab}5"].font = Font(name="Arial", italic=True)
ws[f"{lab}7"] = "Summary"; ws[f"{lab}7"].font = bold
for c, t in zip((lab, med, mn, big, mod), ["pair", "median", "min", "frames < big", "frames < moderate"]): ws[f"{c}8"] = t; ws[f"{c}8"].font = bold
for k, c in enumerate(cols):
    col = L(4 + k); r = 9 + k
    ws[f"{lab}{r}"] = c; ws[f"{med}{r}"] = f"=MEDIAN({col}2:{col}{n})"; ws[f"{mn}{r}"] = f"=MIN({col}2:{col}{n})"
    ws[f"{big}{r}"] = f'=COUNTIF({col}2:{col}{n},"<"&{thr_big})'; ws[f"{mod}{r}"] = f'=COUNTIF({col}2:{col}{n},"<"&{thr_mod})'
    for cc in (lab, med, mn, big, mod): ws[f"{cc}{r}"].font = reg
    ws[f"{med}{r}"].number_format = ws[f"{mn}{r}"].number_format = "0.000"
ws[f"{lab}{10 + len(cols)}"] = "IoU at half resolution (2016x1512); sets: masks/ (baseline), masks_sam2/, masks_sam3/ on OneDrive, out/color_masks_full (repo); compare_iou.py"
ws[f"{lab}{10 + len(cols)}"].font = Font(name="Arial", italic=True)
ws.column_dimensions[lab].width = 30
for cc in (med, mn, big, mod): ws.column_dimensions[cc].width = 16
wb.save("out/bench/iou_pairs.xlsx")

low = df[(df[cols] < BIG).any(axis=1)]
print("big-disagreement frames (<%.2f):" % BIG, list(low.image), "| moderate (<%.2f):" % MOD, int((df[cols] < MOD).any(axis=1).sum()))
sets = {"baseline (old pipeline)": f"{D}/masks", "SAM 2.1": f"{D}/masks_sam2", "SAM 3": f"{D}/masks_sam3", "box + colour key": "out/color_masks_full"}
tiles = []
for _, r in low.iterrows():
    name = r["image"]; rgb = C.load_rgb(f"{D}/stacked/{name}.tif")
    ms = {k: C.load_mask(f"{v}/{name}.png") for k, v in sets.items() if os.path.exists(f"{v}/{name}.png")}
    b, crop, _ = viz.crop_to(rgb, list(ms.values()), pad=0.25)
    bad = [c for c in cols if r[c] < BIG]
    tiles.append(viz.label(crop, name + "  low: " + ", ".join(f"{c} {r[c]:.2f}" for c in bad)[:60]))
    for k, m in ms.items():
        tiles.append(viz.label(viz.overlay(crop, m[b[1]:b[3], b[0]:b[2]], (255, 0, 0), alpha=0.3, thick=2), k))
ncol = 1 + len(sets)
Image.fromarray(viz.grid(tiles, ncol, 440)).save("out/bench/compare_big_disagreement_overlay.jpg", quality=90)
print("written: out/bench/iou_pairs.xlsx, out/bench/compare_big_disagreement_overlay.jpg")
